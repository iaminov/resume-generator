"""Tests for DOCX to markdown conversion.

This sits at the front of /parse-resumes, so every profile is built on its
output. Fixtures are constructed with python-docx at test time rather than
committed as binaries, which keeps them readable and diffable.
"""
import pytest
from docx import Document

from docx_to_md import docx_to_md, para_to_md


def build(tmp_path, build_fn, name="sample.docx"):
    doc = Document()
    build_fn(doc)
    path = tmp_path / name
    doc.save(str(path))
    return path


class TestParagraphConversion:
    def test_plain_text_passes_through(self, tmp_path):
        path = build(tmp_path, lambda d: d.add_paragraph("Just a sentence."))
        assert "Just a sentence." in docx_to_md(path)

    def test_headings_become_hashes(self, tmp_path):
        def content(d):
            d.add_paragraph("Experience", style="Heading 1")
            d.add_paragraph("Deeper", style="Heading 2")
        md = docx_to_md(build(tmp_path, content))
        assert "# Experience" in md
        assert "## Deeper" in md

    def test_bullets_become_dashes(self, tmp_path):
        def content(d):
            d.add_paragraph("Built a thing.", style="List Bullet")
        assert "- Built a thing." in docx_to_md(build(tmp_path, content))

    def test_title_becomes_top_level_heading(self, tmp_path):
        path = build(tmp_path, lambda d: d.add_paragraph("Jane Doe", style="Title"))
        assert "# Jane Doe" in docx_to_md(path)

    def test_empty_paragraphs_produce_nothing(self, tmp_path):
        doc = Document()
        para = doc.add_paragraph("   ")
        assert para_to_md(para) == ""

    def test_whitespace_is_trimmed(self, tmp_path):
        doc = Document()
        assert para_to_md(doc.add_paragraph("  padded  ")) == "padded"


class TestTables:
    def test_table_cells_are_extracted(self, tmp_path):
        # Resumes often use tables purely for layout, so their content must
        # survive conversion or the profile silently loses it.
        def content(d):
            table = d.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Acme Corp"
            table.cell(0, 1).text = "2020 - 2024"
            table.cell(1, 0).text = "Senior Engineer"
            table.cell(1, 1).text = "Springfield"
        md = docx_to_md(build(tmp_path, content))
        for expected in ("Acme Corp", "2020 - 2024", "Senior Engineer", "Springfield"):
            assert expected in md, expected


class TestWholeDocument:
    def test_a_resume_shaped_document_converts(self, tmp_path):
        def content(d):
            d.add_paragraph("Jane Doe", style="Title")
            d.add_paragraph("Experience", style="Heading 1")
            d.add_paragraph("Senior Engineer, Acme Corp")
            d.add_paragraph("Built the payments service.", style="List Bullet")
            d.add_paragraph("Reduced latency on the hot path.", style="List Bullet")
        md = docx_to_md(build(tmp_path, content))
        assert "# Jane Doe" in md
        assert "# Experience" in md
        assert "- Built the payments service." in md
        assert md.endswith("\n")

    def test_ordering_is_preserved(self, tmp_path):
        def content(d):
            for word in ("first", "second", "third"):
                d.add_paragraph(word)
        md = docx_to_md(build(tmp_path, content))
        assert md.index("first") < md.index("second") < md.index("third")

    def test_empty_document_does_not_crash(self, tmp_path):
        path = build(tmp_path, lambda d: None, name="empty.docx")
        assert isinstance(docx_to_md(path), str)

    def test_unicode_survives(self, tmp_path):
        # Resumes carry em-dashes, accents, and smart quotes constantly.
        text = "Café — naïve “quoted” résumé"
        path = build(tmp_path, lambda d: d.add_paragraph(text))
        md = docx_to_md(path)
        for ch in ("Café", "—", "naïve", "résumé"):
            assert ch in md, ch

    def test_missing_file_raises(self, tmp_path):
        with pytest.raises(Exception):
            docx_to_md(tmp_path / "does-not-exist.docx")
