"""Tests for --template handling.

A template supplies styles, not content. Both failure modes here were real
before the loader existed: leftover body content silently prepended itself to
every generated document, and a template lacking "List Bullet" died with a bare
KeyError from deep inside python-docx.
"""
import pytest
from docx import Document

from layout import REQUIRED_STYLES, open_base_document
import generate_resume
import generate_cover_letter

REPO_TEMPLATE = "templates/default.docx"


@pytest.fixture
def template_with_content(tmp_path):
    doc = Document()
    doc.add_paragraph("Leftover heading")
    doc.add_paragraph("Leftover body text")
    path = tmp_path / "with_content.docx"
    doc.save(str(path))
    return path


@pytest.fixture
def template_missing_bullet(tmp_path):
    doc = Document()
    for style in list(doc.styles):
        if style.name == "List Bullet":
            doc.styles.element.remove(style.element)
    path = tmp_path / "no_bullet.docx"
    doc.save(str(path))
    return path


class TestOpenBaseDocument:
    def test_no_template_gives_a_blank_document(self):
        doc, dropped = open_base_document(None)
        assert dropped == []
        assert len(doc.paragraphs) == 0

    def test_missing_file_is_a_clear_error(self, tmp_path):
        with pytest.raises(ValueError, match="template not found"):
            open_base_document(tmp_path / "absent.docx")

    def test_missing_required_style_names_the_style(self, template_missing_bullet):
        with pytest.raises(ValueError, match="List Bullet"):
            open_base_document(template_missing_bullet)

    def test_body_content_is_stripped_and_reported(self, template_with_content):
        doc, dropped = open_base_document(template_with_content)
        assert len(doc.paragraphs) == 0
        assert "Leftover heading" in dropped
        assert "Leftover body text" in dropped

    def test_tables_are_stripped_too(self, tmp_path):
        doc = Document()
        doc.add_table(rows=2, cols=2)
        path = tmp_path / "with_table.docx"
        doc.save(str(path))
        opened, _ = open_base_document(path)
        assert len(opened.tables) == 0


class TestShippedTemplate:
    def test_it_exists_and_is_empty(self):
        doc = Document(REPO_TEMPLATE)
        assert len(doc.paragraphs) == 0, "a template must carry styles, not content"
        assert len(doc.tables) == 0

    def test_it_defines_every_required_style(self):
        available = {s.name for s in Document(REPO_TEMPLATE).styles}
        for style in REQUIRED_STYLES:
            assert style in available, style

    def test_it_loads_without_complaint(self):
        _, dropped = open_base_document(REPO_TEMPLATE)
        assert dropped == []


class TestGeneratorsAcceptIt:
    def test_resume_generates_from_the_template(self, tmp_path, short_resume):
        out = tmp_path / "resume.docx"
        _, _, dropped = generate_resume.generate_resume(
            short_resume, out, template_path=REPO_TEMPLATE
        )
        assert out.exists() and dropped == []
        assert Document(str(out)).paragraphs[0].text.strip() != ""

    def test_cover_letter_generates_from_the_template(self, tmp_path, cover_letter):
        out = tmp_path / "letter.docx"
        dropped = generate_cover_letter.generate_cover_letter(
            cover_letter, out, template_path=REPO_TEMPLATE
        )
        assert out.exists() and dropped == []

    def test_leftover_content_never_reaches_the_output(self, tmp_path, short_resume, template_with_content):
        out = tmp_path / "resume.docx"
        _, _, dropped = generate_resume.generate_resume(
            short_resume, out, template_path=template_with_content
        )
        text = "\n".join(p.text for p in Document(str(out)).paragraphs)
        assert "Leftover" not in text
        assert dropped, "the drop should be reported, not silent"
