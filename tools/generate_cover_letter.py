#!/usr/bin/env python3
"""Generate a .docx cover letter from a JSON content file.

This script is a GENERIC formatter. It takes structured JSON describing the
letter (already written by Claude) and produces a formatted Word document. It
does NOT contain any person-specific content -- everything comes from the JSON.

Usage:
    python3 tools/generate_cover_letter.py content.json output.docx
    python3 tools/generate_cover_letter.py content.json output.docx --template templates/letter.docx

A cover letter is a business letter, not a resume: block format, flush left, no
indentation, a blank line between paragraphs, and generous margins. It should
fit one page essentially always, so the script warns rather than reflowing when
it does not -- an overlong letter needs cutting, not tighter leading.

JSON input format:
{
  "name": "Full Name",
  "contact": {
    "email": "...", "phone": "...", "location": "...",
    "linkedin": "...", "github": "...", "portfolio": "..."
  },
  "date": "August 2, 2026",
  "recipient": {
    "name": "Ms. Jane Smith",          # optional
    "title": "Engineering Manager",     # optional
    "company": "Acme Corp",
    "location": "New York, NY"          # optional
  },
  "salutation": "Dear Ms. Smith,",
  "body": [
    "Opening paragraph...",
    "Evidence paragraph...",
    "Closing paragraph..."
  ],
  "closing": "Sincerely,",
  "signature": "Full Name"              # defaults to "name"
}

Only "name" and "body" are required. Missing optional blocks are skipped.
"""
import argparse
import json
import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Reuse the resume generator's render-based page check so both tools agree on
# what "one page" means. tools/ is on sys.path when these scripts are run.
try:
    from generate_resume import verify_layout
except ImportError:  # pragma: no cover - only when imported as a module
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from generate_resume import verify_layout


# Letter formatting. Body runs a point larger than the resume's: a letter is
# read as prose rather than scanned, and 11pt over 1in margins is the
# conventional business-letter measure.
FONT_NAME = "Calibri"
NAME_SIZE = Pt(16)
BODY_SIZE = Pt(11)
CONTACT_SIZE = Pt(10)
MARGIN = Inches(1.0)
PARAGRAPH_SPACING = Pt(10)
COLOR_BODY = RGBColor(0x00, 0x00, 0x00)
COLOR_META = RGBColor(0x55, 0x55, 0x55)


def set_font(run, size=BODY_SIZE, bold=False, color=COLOR_BODY):
    """Apply font formatting to a run."""
    run.font.name = FONT_NAME
    run.font.size = size
    run.font.bold = bold
    run.font.color.rgb = color


def add_paragraph(doc, text="", size=BODY_SIZE, bold=False, color=COLOR_BODY,
                  space_after=0, align=None):
    """Add a single formatted paragraph."""
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after = Pt(space_after)
    if align is not None:
        para.alignment = align
    if text:
        run = para.add_run(text)
        set_font(run, size=size, bold=bold, color=color)
    return para


def build_letterhead(doc, data):
    """Sender name and contact details, centred to echo the resume header."""
    add_paragraph(
        doc, data["name"].upper(), size=NAME_SIZE, bold=True,
        space_after=2, align=WD_ALIGN_PARAGRAPH.CENTER,
    )

    contact = data.get("contact") or {}
    line = [contact[f] for f in ("email", "phone", "location") if contact.get(f)]
    links = [contact[f] for f in ("linkedin", "github", "portfolio") if contact.get(f)]
    if line:
        add_paragraph(
            doc, " | ".join(line), size=CONTACT_SIZE, color=COLOR_META,
            space_after=2, align=WD_ALIGN_PARAGRAPH.CENTER,
        )
    if links:
        add_paragraph(
            doc, " | ".join(links), size=CONTACT_SIZE, color=COLOR_META,
            space_after=18, align=WD_ALIGN_PARAGRAPH.CENTER,
        )
    elif line:
        doc.paragraphs[-1].paragraph_format.space_after = Pt(18)


def build_date(doc, data):
    """Letter date, flush left. Defaults to today in long form."""
    text = (data.get("date") or "").strip()
    if not text:
        text = date.today().strftime("%B %-d, %Y") if sys.platform != "win32" \
            else date.today().strftime("%B %d, %Y").replace(" 0", " ")
    add_paragraph(doc, text, space_after=14)


def build_recipient(doc, data):
    """Recipient block: name, title, company, location -- one line each."""
    recipient = data.get("recipient") or {}
    lines = [recipient.get(f) for f in ("name", "title", "company", "location")]
    lines = [line for line in lines if line]
    if not lines:
        return
    for i, line in enumerate(lines):
        add_paragraph(doc, line, space_after=14 if i == len(lines) - 1 else 0)


def build_body(doc, data):
    """Salutation, body paragraphs, closing, and signature."""
    salutation = (data.get("salutation") or "Dear Hiring Manager,").strip()
    add_paragraph(doc, salutation, space_after=PARAGRAPH_SPACING.pt)

    body = data.get("body") or []
    if isinstance(body, str):
        body = [p.strip() for p in body.split("\n\n") if p.strip()]
    for paragraph in body:
        add_paragraph(doc, paragraph.strip(), space_after=PARAGRAPH_SPACING.pt)

    closing = (data.get("closing") or "Sincerely,").strip()
    add_paragraph(doc, closing, space_after=24)
    add_paragraph(doc, data.get("signature") or data["name"])


def generate_cover_letter(data, output_path, template_path=None):
    """Generate the complete .docx cover letter from JSON data."""
    if template_path and Path(template_path).exists():
        doc = Document(template_path)
    else:
        doc = Document()

    for section in doc.sections:
        section.top_margin = MARGIN
        section.bottom_margin = MARGIN
        section.left_margin = MARGIN
        section.right_margin = MARGIN

    style = doc.styles["Normal"]
    style.font.name = FONT_NAME
    style.font.size = BODY_SIZE
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)

    build_letterhead(doc, data)
    build_date(doc, data)
    build_recipient(doc, data)
    build_body(doc, data)

    doc.save(str(output_path))


def word_count(data):
    """Words in the body only -- the part that should be kept short."""
    body = data.get("body") or []
    if isinstance(body, str):
        body = [body]
    return sum(len(p.split()) for p in body)


def main():
    parser = argparse.ArgumentParser(
        description="Generate a .docx cover letter from JSON content"
    )
    parser.add_argument("content", help="Path to JSON file with letter content")
    parser.add_argument("output", help="Output .docx file path")
    parser.add_argument("--template", help="Optional .docx template file")
    args = parser.parse_args()

    content_path = Path(args.content)
    if not content_path.exists():
        print(f"Error: {content_path} not found", file=sys.stderr)
        sys.exit(1)

    try:
        data = json.loads(content_path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as e:
        print(f"Error: Invalid JSON in {content_path}: {e}", file=sys.stderr)
        sys.exit(1)

    for field in ("name", "body"):
        if not data.get(field):
            print(f"Error: JSON must include '{field}'", file=sys.stderr)
            sys.exit(1)

    output_path = Path(args.output)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    generate_cover_letter(data, output_path, args.template)

    words = word_count(data)
    pages, _ = verify_layout(output_path)
    warnings = []
    if pages and pages > 1:
        warnings.append(
            f"letter runs to {pages} pages -- cut it to one; a cover letter is "
            f"never improved by a second page"
        )
    if words > 400:
        warnings.append(
            f"body is {words} words; 250-400 is the readable range for a cover letter"
        )

    result = {
        "status": "generated",
        "output": str(output_path),
        "body_paragraphs": len(data.get("body") or []),
        "body_words": words,
        "pages": pages,
        "warnings": warnings,
    }
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
