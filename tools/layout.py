#!/usr/bin/env python3
"""Text metrics and page verification shared by the document generators.

Two jobs, both independent of what kind of document is being produced:

1. **Text metrics** -- measuring how many lines a string occupies at a given
   font size and column width. python-docx cannot paginate, so a generator that
   wants to predict height before writing the file has to compute it.

2. **Page verification** -- rendering a finished .docx and reporting how many
   pages it really came to, and how full the last one is. Estimation is
   approximate; when the number matters, render and look.

This lives apart from the generators so that neither has to import the other.
Both are imported as siblings, which works because running tools/<script>.py
puts tools/ on sys.path.
"""
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document

# Approximate line-height multiplier for Calibri at single spacing.
LINE_HEIGHT = 1.22

REQUIRED_STYLES = ("Normal", "List Bullet")


def open_base_document(template_path=None):
    """Open a template as the starting document, or a blank one.

    A template supplies STYLES, not content. python-docx opens a .docx whole, so
    any body content in the template would be prepended to every document
    generated from it -- silently, and on every run. Strip it.

    Templates are also not interchangeable: the generators reference "List
    Bullet" by name, and a .docx saved out of Word that never used a bulleted
    list will not define it. Catch that here with a message that says what to
    fix, rather than letting a bare KeyError surface later.
    """
    if not template_path:
        return Document(), []

    path = Path(template_path)
    if not path.exists():
        raise ValueError(f"template not found: {path}")

    doc = Document(str(path))

    available = {s.name for s in doc.styles}
    missing = [s for s in REQUIRED_STYLES if s not in available]
    if missing:
        raise ValueError(
            f"template {path.name} is missing required style(s): "
            f"{', '.join(missing)}. Add them in Word (apply the style once, "
            f"then delete the text) or start from templates/default.docx."
        )

    removed = [p.text for p in doc.paragraphs if p.text.strip()]
    for para in list(doc.paragraphs):
        para._element.getparent().remove(para._element)
    for table in list(doc.tables):
        table._element.getparent().remove(table._element)

    return doc, removed


# Which font file the metrics actually resolved to, or None if none was found
# and the crude character-width heuristic is in use. The estimate is calibrated
# against Calibri; anything else shifts its accuracy, so this is worth surfacing
# rather than leaving invisible.
FONT_SOURCE = None


def _font_loader():
    """Return a callable (size_pt, bold) -> PIL font, or None if unavailable."""
    try:
        from PIL import ImageFont
    except ImportError:
        return None

    candidates = {
        False: ["calibri.ttf", "Calibri.ttf", "arial.ttf", "DejaVuSans.ttf"],
        True: ["calibrib.ttf", "Calibrib.ttf", "arialbd.ttf", "DejaVuSans-Bold.ttf"],
    }
    search_dirs = [
        Path('C:\\Windows\\Fonts'),
        Path("/usr/share/fonts/truetype/dejavu"),
        Path("/Library/Fonts"),
        Path("/System/Library/Fonts/Supplemental"),
    ]
    resolved = {}
    for bold, names in candidates.items():
        for directory in search_dirs:
            for name in names:
                candidate = directory / name
                if candidate.exists():
                    resolved[bold] = str(candidate)
                    break
            if bold in resolved:
                break
    if not resolved:
        return None

    global FONT_SOURCE
    FONT_SOURCE = Path(resolved[False]).name if False in resolved else         Path(next(iter(resolved.values()))).name

    cache = {}

    def load(size_pt, bold=False):
        path = resolved.get(bold) or next(iter(resolved.values()))
        # Render at 4x for sub-point metric precision, then scale back.
        key = (path, round(size_pt * 4))
        if key not in cache:
            try:
                cache[key] = ImageFont.truetype(path, int(size_pt * 4))
            except OSError:
                return None
        return cache[key]

    return load


_LOAD_FONT = _font_loader()


def text_width_pt(text, size_pt, bold=False):
    """Width of `text` in points, via font metrics when available."""
    if _LOAD_FONT is not None:
        font = _LOAD_FONT(size_pt, bold)
        if font is not None:
            return font.getlength(text) / 4.0
    # Fallback: Calibri averages roughly 0.48 em per character in prose.
    return len(text) * size_pt * 0.48


def wrapped_lines(text, size_pt, avail_pt, bold=False, indent_pt=0.0):
    """Number of rendered lines `text` occupies, by greedy word wrap."""
    avail = max(avail_pt - indent_pt, 1.0)
    words = text.split()
    if not words:
        return 1
    space = text_width_pt(" ", size_pt, bold)
    lines, current = 1, 0.0
    for word in words:
        width = text_width_pt(word, size_pt, bold)
        if current and current + space + width > avail:
            lines += 1
            current = width
        else:
            current += (space if current else 0.0) + width
    return lines


def verify_layout(docx_path):
    """Return (page_count, last_page_fill) for a generated .docx, or (None, None).

    `last_page_fill` is the final page's line count as a fraction of the fullest
    page, so a trailing page holding two lines of a forty-line layout scores
    about 0.05. Converts via LibreOffice when installed; optional by design, so
    callers treat None as "unknown" rather than as an error.
    """
    pdf_bytes = _render_pdf(docx_path)
    if pdf_bytes is None:
        return None, None
    pages = len(re.findall(rb"/Type\s*/Page[^s]", pdf_bytes)) or None
    if pages is None:
        return None, None
    if pages == 1:
        return 1, 1.0

    try:
        import io

        import pdfplumber

        with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
            counts = [
                len((page.extract_text() or "").splitlines()) for page in pdf.pages
            ]
    except Exception:
        return pages, None

    fullest = max(counts) if counts else 0
    if not fullest:
        return pages, None
    return pages, counts[-1] / fullest


def _render_pdf(docx_path):
    """Render a .docx to PDF bytes via LibreOffice, or None if unavailable.

    LibreOffice rather than Word: it runs headless on every platform, needs no
    licence, and will not block on a modal dialog the way COM automation can.
    Its layout is very close to Word's but not identical, so a document sitting
    exactly on a page boundary may still differ by a line.
    """
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice is None:
        for candidate in (
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
            "/Applications/LibreOffice.app/Contents/MacOS/soffice",
        ):
            if Path(candidate).exists():
                soffice = candidate
                break
    if soffice is None:
        return None

    docx_path = Path(docx_path)
    with tempfile.TemporaryDirectory() as tmp:
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir", tmp,
                 str(docx_path)],
                check=True, capture_output=True, timeout=120,
            )
        except (subprocess.SubprocessError, OSError):
            return None
        pdf = Path(tmp) / (docx_path.stem + ".pdf")
        if not pdf.exists():
            return None
        return pdf.read_bytes()


def verify_page_count(docx_path):
    """Return the true page count of a generated .docx, or None if unavailable."""
    return verify_layout(docx_path)[0]


def metrics_source():
    """Describe how text widths are being measured, for diagnostics.

    Returns the font filename when real metrics are in use, or "heuristic" when
    no usable font was found and widths fall back to an average-character-width
    approximation. The calibration constant in generate_resume.py was measured
    against Calibri, so a different source means a different accuracy.
    """
    if _LOAD_FONT is None or FONT_SOURCE is None:
        return "heuristic"
    return FONT_SOURCE
